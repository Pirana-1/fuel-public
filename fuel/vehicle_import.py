import re
from decimal import Decimal

from django.core.exceptions import ValidationError
from django.db import transaction
from django.db.models import Exists, OuterRef
from openpyxl import load_workbook

from .models import FuelMovement, Region, Vehicle


EXPECTED_HEADERS = (
    "Plaka",
    "Araç Anahtarı",
    "Açıklama",
    "Makina Kodu",
    "saat / km / taseron",
)
SOURCE_TYPE_VALUES = {
    "km": (Vehicle.MeterType.KM, False),
    "saat": (Vehicle.MeterType.HOUR, False),
    "taseron": (Vehicle.MeterType.NONE, True),
    "taşeron": (Vehicle.MeterType.NONE, True),
}
UNASSIGNED_REGION = {
    "region_name": "Bölgesi belirlenmedi",
    "region_code": "BELIRSIZ",
}
NUMBERED_REGION_RE = re.compile(
    r"(?P<yht>\bYHT\s+)?(?P<number>[1-5])\s*\.\s*BÖLGE\b",
    re.IGNORECASE,
)
HEADQUARTERS_RE = re.compile(
    r"(?P<yht>\bYHT\s+)?(?:GENEL\s+MÜDÜRLÜK|GEN\.\s*MÜD\.)",
    re.IGNORECASE,
)


def _cell_text(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, (int, Decimal)):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _tag_text(value):
    if isinstance(value, float) and not value.is_integer():
        raise ValueError("Araç anahtarı tam sayı veya metin olmalıdır.")
    return _cell_text(value).strip().upper()


def _source_type_text(value):
    source_type = SOURCE_TYPE_VALUES.get(_cell_text(value).strip().casefold())
    if source_type is None:
        raise ValueError("Sayaç sütununda yalnızca km, saat veya taşeron yazmalıdır.")
    return source_type


def _organization_around(description, match):
    parts = []
    if match.groupdict().get("yht"):
        parts.append("YHT")
    before = description[: match.start()].strip(" -·")
    after = description[match.end() :].strip(" -·")
    if before:
        parts.append(before)
    if after:
        parts.append(after)
    return " · ".join(parts)


def parse_description(description):
    description = " ".join(description.split())
    numbered = NUMBERED_REGION_RE.search(description)
    if numbered:
        number = numbered.group("number")
        return {
            "region_name": f"{number}. Bölge",
            "region_code": f"{number}.BOLGE",
            "organization": _organization_around(description, numbered),
            "region_was_inferred": False,
        }

    headquarters = HEADQUARTERS_RE.search(description)
    if headquarters:
        return {
            "region_name": "Genel Müdürlük",
            "region_code": "GENEL-MUD",
            "organization": _organization_around(description, headquarters),
            "region_was_inferred": False,
        }

    return {
        **UNASSIGNED_REGION,
        "organization": description,
        "region_was_inferred": True,
    }


def parse_vehicle_workbook(workbook_file):
    try:
        workbook = load_workbook(workbook_file, read_only=True, data_only=True)
    except Exception as error:
        raise ValidationError("Excel dosyası okunamadı veya geçerli bir .xlsx dosyası değil.") from error

    try:
        worksheet = workbook.active
        first_row = next(
            worksheet.iter_rows(min_row=1, max_row=1, values_only=True),
            None,
        )
        headers = tuple(
            _cell_text(value) for value in (first_row or ())[: len(EXPECTED_HEADERS)]
        )
        if headers != EXPECTED_HEADERS:
            expected = ", ".join(EXPECTED_HEADERS)
            raise ValidationError(f"Excel başlıkları şu sırada olmalıdır: {expected}.")

        rows = []
        errors = []
        codes = {}
        tags = {}
        for row_number, values in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            values = list(values[: len(EXPECTED_HEADERS)]) + [None] * max(
                0, len(EXPECTED_HEADERS) - len(values)
            )
            if not any(value not in (None, "") for value in values):
                continue

            plate = _cell_text(values[0]).strip().upper()
            machine_code = _cell_text(values[3]).strip()
            description = _cell_text(values[2]).strip()
            try:
                tag_code = _tag_text(values[1])
            except ValueError as error:
                errors.append(f"{row_number}. satır: {error}")
                continue
            try:
                meter_type, is_contractor = _source_type_text(values[4])
            except ValueError as error:
                errors.append(f"{row_number}. satır: {error}")
                continue

            code = plate or machine_code.upper()
            name = machine_code or plate
            row_errors = []
            if not code:
                row_errors.append("plaka veya makina kodu bulunmalıdır")
            if not tag_code:
                row_errors.append("araç anahtarı boş olamaz")
            if not description:
                row_errors.append("açıklama boş olamaz")
            if len(code) > 60:
                row_errors.append("araç kodu/plaka 60 karakteri aşamaz")
            if len(name) > 160:
                row_errors.append("makina kodu 160 karakteri aşamaz")
            if len(tag_code) > 20:
                row_errors.append("araç anahtarı 20 karakteri aşamaz")
            if row_errors:
                errors.append(f"{row_number}. satır: " + "; ".join(row_errors) + ".")
                continue

            parsed_description = parse_description(description)
            if len(parsed_description["organization"]) > 160:
                errors.append(
                    f"{row_number}. satır: açıklama bilgisi 160 karakteri aşamaz."
                )
                continue

            if code in codes:
                errors.append(
                    f"{row_number}. satır: {code} araç kodu {codes[code]}. satırda da kullanılmış."
                )
                continue
            if tag_code in tags:
                errors.append(
                    f"{row_number}. satır: {tag_code} araç anahtarı {tags[tag_code]}. satırda da kullanılmış."
                )
                continue
            codes[code] = row_number
            tags[tag_code] = row_number
            rows.append(
                {
                    "source_row": row_number,
                    "code": code,
                    "tag_code": tag_code,
                    "name": name,
                    "organization": parsed_description["organization"],
                    "region_name": parsed_description["region_name"],
                    "region_code": parsed_description["region_code"],
                    "region_was_inferred": parsed_description["region_was_inferred"],
                    "meter_type": meter_type,
                    "meter_type_label": dict(Vehicle.MeterType.choices)[meter_type],
                    "is_contractor": is_contractor,
                    "vehicle_class_label": (
                        "Taşeron" if is_contractor else "Kurum aracı"
                    ),
                }
            )

        if errors:
            raise ValidationError(errors)
        if not rows:
            raise ValidationError("Excel dosyasında içe aktarılabilecek araç bulunamadı.")
        return {
            "sheet_name": worksheet.title,
            "rows": rows,
            "row_count": len(rows),
            "unassigned_region_count": sum(
                row["region_was_inferred"] for row in rows
            ),
        }
    finally:
        workbook.close()


def assess_vehicle_rows(rows):
    existing_vehicles = list(
        Vehicle.objects.select_related("region").annotate(
            has_fuel_movements=Exists(
                FuelMovement.objects.filter(vehicle_id=OuterRef("pk"))
            )
        )
    )
    by_code = {vehicle.code.upper(): vehicle for vehicle in existing_vehicles}
    by_tag = {
        vehicle.tag_code.upper(): vehicle
        for vehicle in existing_vehicles
        if vehicle.tag_code
    }
    create_count = 0
    update_count = 0
    skip_count = 0
    conflicts = []
    warnings = []
    protected_count = 0
    for row in rows:
        code_vehicle = by_code.get(row["code"].upper())
        tag_vehicle = by_tag.get(row["tag_code"].upper())
        if code_vehicle is None and tag_vehicle is None:
            create_count += 1
        elif (
            code_vehicle is not None
            and tag_vehicle is not None
            and code_vehicle.pk == tag_vehicle.pk
        ):
            meter_type_changed = code_vehicle.meter_type != row["meter_type"]
            contractor_changed = (
                code_vehicle.is_contractor != row["is_contractor"]
            )
            if not meter_type_changed and not contractor_changed:
                skip_count += 1
            else:
                if contractor_changed or not code_vehicle.has_fuel_movements:
                    update_count += 1
            if meter_type_changed and code_vehicle.has_fuel_movements:
                suffix = (
                    " Taşeron sınıfı yine güncellenecek."
                    if contractor_changed
                    else ""
                )
                warnings.append(
                    f"{row['source_row']}. satır: {row['code']} aracının sayaç türü "
                    "geçmiş yakıt hareketleri nedeniyle güncellenmeden bırakılacak."
                    f"{suffix}"
                )
                protected_count += 1
        elif code_vehicle is not None:
            conflicts.append(
                f"{row['source_row']}. satır: {row['code']} kodlu araç farklı bir RFID koduyla zaten kayıtlı."
            )
        else:
            conflicts.append(
                f"{row['source_row']}. satır: {row['tag_code']} RFID kodu {tag_vehicle.code} aracında zaten kayıtlı."
            )
    return {
        "create_count": create_count,
        "update_count": update_count,
        "skip_count": skip_count,
        "import_count": create_count + update_count,
        "conflicts": conflicts,
        "warnings": warnings,
        "protected_count": protected_count,
    }


@transaction.atomic
def import_vehicle_rows(rows):
    if any(row.get("meter_type") not in Vehicle.MeterType.values for row in rows):
        raise ValidationError("Excel satırlarında geçerli bir sayaç türü bulunmalıdır.")
    if any(not isinstance(row.get("is_contractor"), bool) for row in rows):
        raise ValidationError("Excel satırlarında araç sınıfı bulunmalıdır.")

    assessment = assess_vehicle_rows(rows)
    if assessment["conflicts"]:
        raise ValidationError(assessment["conflicts"])

    regions = {}
    for row in rows:
        region_code = row["region_code"]
        if region_code not in regions:
            region, _ = Region.objects.get_or_create(
                code=region_code,
                defaults={"name": row["region_name"], "is_active": True},
            )
            regions[region_code] = region

    existing_vehicles = list(
        Vehicle.objects.annotate(
            has_fuel_movements=Exists(
                FuelMovement.objects.filter(vehicle_id=OuterRef("pk"))
            )
        )
    )
    existing_by_code = {vehicle.code.upper(): vehicle for vehicle in existing_vehicles}
    vehicles_to_update = []
    for row in rows:
        vehicle = existing_by_code.get(row["code"].upper())
        if not (
            vehicle is not None
            and vehicle.tag_code
            and vehicle.tag_code.upper() == row["tag_code"].upper()
        ):
            continue
        changed = False
        if vehicle.is_contractor != row["is_contractor"]:
            vehicle.is_contractor = row["is_contractor"]
            changed = True
        if (
            not vehicle.has_fuel_movements
            and vehicle.meter_type != row["meter_type"]
        ):
            vehicle.meter_type = row["meter_type"]
            changed = True
        if changed:
            vehicles_to_update.append(vehicle)
    if vehicles_to_update:
        Vehicle.objects.bulk_update(
            vehicles_to_update,
            ["meter_type", "is_contractor"],
            batch_size=500,
        )

    existing_pairs = {
        (vehicle.code.upper(), vehicle.tag_code.upper())
        for vehicle in Vehicle.objects.exclude(tag_code__isnull=True)
    }
    vehicles = [
        Vehicle(
            code=row["code"],
            tag_code=row["tag_code"],
            name=row["name"],
            organization=row["organization"],
            region=regions[row["region_code"]],
            meter_type=row["meter_type"],
            is_contractor=row["is_contractor"],
            is_active=True,
        )
        for row in rows
        if (row["code"].upper(), row["tag_code"].upper()) not in existing_pairs
    ]
    Vehicle.objects.bulk_create(vehicles, batch_size=500)
    return {
        "created_count": len(vehicles),
        "updated_count": len(vehicles_to_update),
        "skipped_count": assessment["skip_count"],
        "protected_count": assessment["protected_count"],
        "unassigned_region_count": sum(
            row["region_was_inferred"] for row in rows
        ),
    }
