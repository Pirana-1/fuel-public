from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.contrib.auth import get_user_model
from django.contrib.auth.models import Group, Permission
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from fuel.models import (
    ContractorFuelPrice,
    FuelMovement,
    Region,
    StorageRegionChange,
    StorageUnit,
    Supplier,
    Vehicle,
    VehicleRegionChange,
)
from fuel.services import create_movement


class FuelViewsTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_superuser(
            username="admin", password="test-password", email="admin@example.test"
        )
        self.region = Region.objects.create(name="1. Bölge", code="B1")
        self.supplier = Supplier.objects.create(name="Test Tedarikçi")
        self.fixed = StorageUnit.objects.create(
            name="Sabit Tank",
            code="SABIT-1",
            kind=StorageUnit.Kind.FIXED_TANK,
            region=self.region,
            capacity_liters=Decimal("10000"),
        )
        self.vehicle = Vehicle.objects.create(
            name="Test Aracı",
            code="06TEST01",
            region=self.region,
            meter_type=Vehicle.MeterType.NONE,
        )
        self.now = timezone.now()
        self.receipt = create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.SUPPLIER_RECEIPT,
            data={
                "occurred_at": self.now,
                "supplier": self.supplier,
                "destination_storage": self.fixed,
                "liters": Decimal("1000"),
                "purchase_unit_price": Decimal("42.5000"),
            },
        )
        self.fueling = create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.VEHICLE_FUELING,
            data={
                "occurred_at": self.now,
                "source_storage": self.fixed,
                "vehicle": self.vehicle,
                "liters": Decimal("125"),
            },
        )

    def test_dashboard_requires_login(self):
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 302)
        self.assertIn(reverse("login"), response.url)

    def test_login_page_contains_only_compact_login_form(self):
        response = self.client.get(reverse("login"))

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "login-form panel")
        self.assertContains(response, "Sisteme giriş")
        self.assertContains(response, "Kullanıcı adı")
        self.assertContains(response, "Parola")
        self.assertNotContains(response, "login-visual")
        self.assertNotContains(response, "Yakıt hareketlerini tek merkezden yönetin")
        self.assertNotContains(response, "Tedarikçi</span>")

    def test_dashboard_displays_stock_and_period_totals(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["total_stock"], Decimal("875"))
        self.assertEqual(response.context["totals"]["received"], Decimal("1000"))
        self.assertEqual(response.context["totals"]["consumed"], Decimal("125"))
        self.assertContains(response, 'style="width: 9%"', html=False)
        self.assertContains(response, "1.000 L")
        self.assertContains(response, "875 L")

    def test_dashboard_period_filters_recent_movements(self):
        old_receipt = create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.SUPPLIER_RECEIPT,
            data={
                "occurred_at": self.now - timedelta(days=40),
                "supplier": self.supplier,
                "destination_storage": self.fixed,
                "liters": Decimal("50"),
                "purchase_unit_price": Decimal("42.0000"),
            },
        )

        self.client.force_login(self.user)
        response = self.client.get(reverse("dashboard"), {"period": "7d"})

        recent_ids = {movement.pk for movement in response.context["recent_movements"]}
        self.assertIn(self.receipt.pk, recent_ids)
        self.assertIn(self.fueling.pk, recent_ids)
        self.assertNotIn(old_receipt.pk, recent_ids)

    def test_movement_list_allows_custom_page_size_and_preserves_it(self):
        for index in range(31):
            create_movement(
                user=self.user,
                movement_type=FuelMovement.MovementType.SUPPLIER_RECEIPT,
                data={
                    "occurred_at": self.now - timedelta(minutes=index + 1),
                    "supplier": self.supplier,
                    "destination_storage": self.fixed,
                    "liters": Decimal("1"),
                    "purchase_unit_price": Decimal("42.0000"),
                },
            )

        self.client.force_login(self.user)
        response = self.client.get(reverse("movement-list"), {"page_size": "15"})

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["page_obj"].paginator.per_page, 15)
        self.assertEqual(response.context["filters"]["page_size"], "15")
        self.assertContains(response, "page_size=15")

    def test_movement_list_rejects_invalid_page_size(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("movement-list"), {"page_size": "999"})

        self.assertEqual(response.context["page_obj"].paginator.per_page, 30)
        self.assertEqual(response.context["filters"]["page_size"], "30")

    def test_movement_list_displays_vehicle_meter_value_and_unit(self):
        self.vehicle.meter_type = Vehicle.MeterType.KM
        self.vehicle.save(update_fields=["meter_type"])
        self.fueling.meter_value = Decimal("124500")
        self.fueling.save(update_fields=["meter_value"])

        self.client.force_login(self.user)
        response = self.client.get(reverse("movement-list"))

        self.assertContains(response, "124.500 km")
        self.assertContains(response, "Sayaç")

    def test_supplier_receipt_form_has_clear_required_choices(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("supplier-receipt-create"))
        form = response.context["form"]

        self.assertTrue(form.fields["supplier"].required)
        self.assertEqual(form.fields["supplier"].empty_label, "Tedarikçi seçin")
        self.assertTrue(form.fields["destination_storage"].required)
        self.assertTrue(form.fields["purchase_unit_price"].required)
        self.assertEqual(
            form.fields["destination_storage"].empty_label,
            "Sabit tank seçin",
        )
        self.assertContains(response, "supplier-receipt-fields")
        self.assertContains(response, "Tedarikçi ve sabit tank doğru mu?")
        self.assertNotContains(response, "Taslak")

    def test_contractor_price_definition_and_fueling_snapshot(self):
        self.client.force_login(self.user)
        effective_from = self.now - timedelta(hours=1)
        response = self.client.post(
            reverse("contractor-price-list"),
            {
                "effective_from": timezone.localtime(effective_from).strftime(
                    "%Y-%m-%dT%H:%M"
                ),
                "reference_unit_price": "42.0000",
                "contractor_unit_price": "44.5000",
                "note": "Ağustos taşeron fiyatı",
            },
        )
        self.assertRedirects(response, reverse("contractor-price-list"))
        self.vehicle.is_contractor = True
        self.vehicle.save(update_fields=["is_contractor"])

        form_page = self.client.get(reverse("vehicle-fueling-create"))
        self.assertNotIn(
            "contractor_reference_unit_price",
            form_page.context["form"].fields,
        )
        self.assertNotIn("contractor_unit_price", form_page.context["form"].fields)
        self.assertNotContains(form_page, 'id="contractor-price-definitions"')
        self.assertNotContains(form_page, "Referans ortalama fiyat")
        self.assertNotContains(form_page, "Taşerona yansıtılacak fiyat")

        response = self.client.post(
            reverse("vehicle-fueling-create"),
            {
                "occurred_at": timezone.localtime(self.now).strftime(
                    "%Y-%m-%dT%H:%M"
                ),
                "source_storage": self.fixed.pk,
                "vehicle": self.vehicle.pk,
                "liters": "10.000",
                "meter_value": "",
                "contractor_reference_unit_price": "1.0000",
                "contractor_unit_price": "2.0000",
                "note": "Taşeron dolumu",
            },
        )

        self.assertRedirects(response, reverse("movement-list"))
        movement = FuelMovement.objects.active().order_by("-pk").first()
        self.assertEqual(
            movement.contractor_reference_unit_price,
            Decimal("42.0000"),
        )
        self.assertEqual(movement.contractor_unit_price, Decimal("44.5000"))

    def test_internal_transfer_form_shows_storage_availability_data(self):
        mobile = StorageUnit.objects.create(
            name="Mobil Tanker",
            code="MOBIL-1",
            kind=StorageUnit.Kind.MOBILE_TANKER,
            region=self.region,
            capacity_liters=Decimal("5000"),
        )
        self.client.force_login(self.user)

        response = self.client.get(reverse("internal-transfer-create"))
        form = response.context["form"]
        storage_status = response.context["storage_status"]

        self.assertTrue(form.fields["source_storage"].required)
        self.assertEqual(
            form.fields["source_storage"].empty_label,
            "Kaynak sabit tank seçin",
        )
        self.assertTrue(form.fields["destination_storage"].required)
        self.assertEqual(
            form.fields["destination_storage"].empty_label,
            "Hedef mobil tanker seçin",
        )
        self.assertEqual(storage_status[str(self.fixed.pk)]["current_stock"], "875.000")
        self.assertEqual(storage_status[str(mobile.pk)]["current_stock"], "0")
        self.assertEqual(storage_status[str(mobile.pk)]["capacity"], "5000.000")
        self.assertEqual(storage_status[str(mobile.pk)]["remaining"], "5000.000")
        self.assertContains(response, 'id="operation-storage-status"')
        self.assertContains(response, 'data-storage-status="source"')
        self.assertContains(response, 'data-storage-status="destination"')

    def test_vehicle_fueling_form_shows_stock_and_meter_context(self):
        self.vehicle.meter_type = Vehicle.MeterType.KM
        self.vehicle.save(update_fields=["meter_type"])
        self.fueling.meter_value = Decimal("152698")
        self.fueling.save(update_fields=["meter_value"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("vehicle-fueling-create"))
        form = response.context["form"]
        storage_status = response.context["storage_status"]
        meter_status = response.context["vehicle_meter_status"]

        self.assertTrue(form.fields["source_storage"].required)
        self.assertEqual(
            form.fields["source_storage"].empty_label,
            "Yakıt veren tank veya tanker seçin",
        )
        self.assertTrue(form.fields["vehicle"].required)
        self.assertEqual(
            form.fields["vehicle"].empty_label,
            "Araç veya makine seçin",
        )
        self.assertEqual(form.fields["note"].widget.attrs["rows"], 3)
        self.assertEqual(storage_status[str(self.fixed.pk)]["current_stock"], "875.000")
        self.assertEqual(meter_status[str(self.vehicle.pk)]["meter_type"], "KM")
        self.assertFalse(meter_status[str(self.vehicle.pk)]["is_contractor"])
        self.assertEqual(meter_status[str(self.vehicle.pk)]["unit"], "km")
        self.assertEqual(meter_status[str(self.vehicle.pk)]["last_value"], "152698.00")
        self.assertContains(response, 'id="operation-storage-status"')
        self.assertContains(response, 'id="vehicle-meter-status"')
        self.assertContains(response, 'data-storage-status="source"')
        self.assertContains(response, 'data-vehicle-meter="true"')

    def test_vehicle_fueling_form_requires_meter_for_metered_vehicle(self):
        self.vehicle.meter_type = Vehicle.MeterType.KM
        self.vehicle.save(update_fields=["meter_type"])
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("vehicle-fueling-create"),
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "source_storage": self.fixed.pk,
                "vehicle": self.vehicle.pk,
                "liters": "10.000",
                "meter_value": "",
                "note": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertFormError(
            response.context["form"],
            "meter_value",
            "Bu araç için kilometre girilmelidir.",
        )

    def test_vehicle_fueling_form_rejects_meter_for_contractor_vehicle(self):
        self.vehicle.is_contractor = True
        self.vehicle.save(update_fields=["is_contractor"])
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("vehicle-fueling-create"),
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "source_storage": self.fixed.pk,
                "vehicle": self.vehicle.pk,
                "liters": "10.000",
                "meter_value": "1000",
                "contractor_reference_unit_price": "42.0000",
                "contractor_unit_price": "44.5000",
                "note": "",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertFormError(
            response.context["form"],
            "meter_value",
            "Taşeron araçlarda sayaç değeri girilmez.",
        )

    def test_negative_stock_is_saved_with_warning_and_visible(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("vehicle-fueling-create"),
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "source_storage": self.fixed.pk,
                "vehicle": self.vehicle.pk,
                "liters": "1000.000",
                "meter_value": "",
                "note": "",
            },
        )

        self.assertEqual(response.status_code, 302)
        self.assertEqual(response["Location"], reverse("movement-list"))
        movement_page = self.client.get(reverse("movement-list"))
        self.assertContains(movement_page, "stoğu eksiye düştü")
        storage_page = self.client.get(reverse("storage-list"))
        self.assertContains(storage_page, "Eksi stok")
        self.assertContains(storage_page, "-125 L")
        export = self.client.get(reverse("report-export-xlsx"))
        workbook = load_workbook(BytesIO(export.content), data_only=True)
        storage_sheet = workbook["Stok Noktaları"]
        storage_headers = [cell.value for cell in storage_sheet[1]]
        self.assertEqual(
            storage_sheet.cell(row=2, column=storage_headers.index("Stok durumu") + 1).value,
            "Eksi stok",
        )

    def test_contractor_report_and_excel_are_separate_and_labeled(self):
        self.vehicle.is_contractor = True
        self.vehicle.save(update_fields=["is_contractor"])
        self.fueling.contractor_reference_unit_price = Decimal("42.0000")
        self.fueling.contractor_unit_price = Decimal("44.5000")
        self.fueling.save(
            update_fields=[
                "contractor_reference_unit_price",
                "contractor_unit_price",
            ]
        )
        self.client.force_login(self.user)

        report = self.client.get(reverse("contractor-reports"))

        self.assertEqual(report.status_code, 200)
        self.assertEqual(report.context["totals"]["consumed"], Decimal("125"))
        self.assertEqual(
            report.context["contractor_stats"]["reference_total"],
            Decimal("5250.0000000"),
        )
        self.assertEqual(
            report.context["contractor_stats"]["charged_total"],
            Decimal("5562.5000000"),
        )
        self.assertContains(report, "Taşeron yakıt dolumları")
        self.assertContains(report, "Taşeron")

        export = self.client.get(reverse("contractor-report-export-xlsx"))
        workbook = load_workbook(BytesIO(export.content), data_only=True)
        self.assertIn("Taşeron Dolumları", workbook.sheetnames)
        contractor_sheet = workbook["Taşeron Dolumları"]
        headers = [cell.value for cell in contractor_sheet[1]]
        self.assertIn("Referans ortalama fiyat", headers)
        self.assertIn("Taşeron birim fiyatı", headers)
        self.assertIn("Toplam fark", headers)
        self.assertEqual(
            contractor_sheet.cell(
                row=2,
                column=headers.index("Referans ortalama fiyat") + 1,
            ).value,
            42,
        )
        self.assertEqual(
            contractor_sheet.cell(
                row=2,
                column=headers.index("Taşeron birim fiyatı") + 1,
            ).value,
            44.5,
        )
        self.assertEqual(contractor_sheet.max_row, 3)
        self.assertIn("Fiyat Tanımları", workbook.sheetnames)

    def test_report_and_excel_display_vehicle_meter_value_and_unit(self):
        self.vehicle.meter_type = Vehicle.MeterType.KM
        self.vehicle.save(update_fields=["meter_type"])
        self.fueling.meter_value = Decimal("124500")
        self.fueling.save(update_fields=["meter_value"])

        self.client.force_login(self.user)
        report = self.client.get(reverse("reports"))

        self.assertContains(report, "124.500 km")
        self.assertContains(report, "Sayaç")
        self.assertContains(report, "İşlem türü")
        self.assertContains(report, "İrsaliye")
        self.assertContains(report, "42,5 TL/L")

        export = self.client.get(reverse("report-export-xlsx"))
        workbook = load_workbook(BytesIO(export.content), data_only=True)
        movements = workbook["Hareketler"]
        movement_columns = {
            cell.value: cell.column - 1 for cell in movements[1]
        }
        fueling_row = next(
            row for row in movements.iter_rows(min_row=2) if row[1].value == "Araca yakıt verme"
        )
        self.assertEqual(fueling_row[movement_columns["Sayaç"]].value, 124500)
        self.assertEqual(fueling_row[movement_columns["Sayaç birimi"]].value, "km")
        self.assertEqual(fueling_row[movement_columns["Sayaç türü"]].value, "Kilometre")
        self.assertEqual(fueling_row[movement_columns["Araç sınıfı"]].value, "Kurum aracı")
        receipt_row = next(
            row
            for row in movements.iter_rows(min_row=2)
            if row[1].value == "Tedarikçiden giriş"
        )
        self.assertEqual(
            receipt_row[movement_columns["Alış birim fiyatı"]].value,
            42.5,
        )
        self.assertEqual(
            receipt_row[movement_columns["Alış toplamı"]].value,
            42500,
        )
        supplier_sheet = workbook["Tedarikçi Özeti"]
        supplier_headers = [cell.value for cell in supplier_sheet[1]]
        self.assertIn("Ağırlıklı ortalama alış fiyatı", supplier_headers)
        self.assertIn("Toplam alış tutarı", supplier_headers)
        self.assertIn("Tedarikçi Alımları", workbook.sheetnames)
        supplier_receipts = workbook["Tedarikçi Alımları"]
        supplier_receipt_headers = [cell.value for cell in supplier_receipts[1]]
        self.assertIn("Alış birim fiyatı", supplier_receipt_headers)
        self.assertIn("Alış toplamı", supplier_receipt_headers)
        self.assertEqual(
            supplier_receipts.cell(
                row=2,
                column=supplier_receipt_headers.index("Alış birim fiyatı") + 1,
            ).value,
            42.5,
        )

    def test_vehicle_filter_only_returns_selected_vehicle(self):
        other_vehicle = Vehicle.objects.create(
            name="Diğer Araç",
            code="06TEST02",
            region=self.region,
            meter_type=Vehicle.MeterType.NONE,
        )
        self.client.force_login(self.user)
        response = self.client.get(reverse("reports"), {"vehicle": other_vehicle.pk})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["consumed"], Decimal("0"))
        self.assertEqual(response.context["vehicle_summary"], [])

    def test_date_filter_uses_inclusive_local_calendar_day(self):
        past = self.now - timedelta(days=40)
        create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.VEHICLE_FUELING,
            data={
                "occurred_at": past,
                "source_storage": self.fixed,
                "vehicle": self.vehicle,
                "liters": Decimal("50"),
            },
        )

        self.client.force_login(self.user)
        response = self.client.get(
            reverse("reports"),
            {"start": past.date().isoformat(), "end": past.date().isoformat()},
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["totals"]["consumed"], Decimal("50"))

    def test_quick_period_overrides_submitted_date_fields(self):
        inside_period = self.now - timedelta(days=6)
        outside_period = self.now - timedelta(days=7)
        for occurred_at, liters in (
            (inside_period, "50"),
            (outside_period, "25"),
        ):
            create_movement(
                user=self.user,
                movement_type=FuelMovement.MovementType.VEHICLE_FUELING,
                data={
                    "occurred_at": occurred_at,
                    "source_storage": self.fixed,
                    "vehicle": self.vehicle,
                    "liters": Decimal(liters),
                },
            )

        self.client.force_login(self.user)
        response = self.client.get(
            reverse("reports"),
            {
                "period": "7d",
                "start": "2000-01-01",
                "end": "2000-01-01",
            },
        )

        expected_start = timezone.localdate() - timedelta(days=6)
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.context["filters"]["period"], "7d")
        self.assertEqual(response.context["filters"]["start"], expected_start.isoformat())
        self.assertEqual(response.context["filters"]["end"], timezone.localdate().isoformat())
        self.assertEqual(response.context["totals"]["consumed"], Decimal("175"))

    def test_primary_pages_render_for_manager(self):
        self.client.force_login(self.user)
        route_names = [
            "dashboard",
            "movement-list",
            "opening-balance-create",
            "supplier-receipt-create",
            "internal-transfer-create",
            "vehicle-fueling-create",
            "stock-adjustment-create",
            "region-list",
            "region-create",
            "supplier-list",
            "supplier-create",
            "storage-list",
            "storage-create",
            "vehicle-list",
            "vehicle-create",
            "reports",
            "user-list",
            "user-create",
        ]

        for route_name in route_names:
            with self.subTest(route_name=route_name):
                response = self.client.get(reverse(route_name))
                self.assertEqual(response.status_code, 200)

        response = self.client.get(reverse("movement-void", args=[self.fueling.pk]))
        self.assertEqual(response.status_code, 200)

    def test_region_list_uses_separate_create_page_and_shows_active_assignments(self):
        StorageUnit.objects.create(
            name="Pasif Tank",
            code="PASIF-TANK",
            kind=StorageUnit.Kind.FIXED_TANK,
            region=self.region,
            capacity_liters=Decimal("1000"),
            is_active=False,
        )
        Vehicle.objects.create(
            name="Pasif Araç",
            code="06PASIF",
            region=self.region,
            meter_type=Vehicle.MeterType.NONE,
            is_active=False,
        )
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("region-list"))
        region = next(item for item in list_page.context["items"] if item == self.region)

        self.assertEqual(list_page.status_code, 200)
        self.assertEqual(region.active_storage_count, 1)
        self.assertEqual(region.active_vehicle_count, 1)
        self.assertContains(list_page, "region-list-panel")
        self.assertContains(list_page, "Aktif tank/tanker")
        self.assertContains(list_page, "Aktif araç/makine")
        self.assertContains(list_page, "Bölge adı veya kodu ara")
        self.assertContains(list_page, f'href="{reverse("region-create")}"')
        self.assertNotContains(list_page, "Yeni bölge</h2>")

        blocked = self.client.post(
            reverse("region-toggle", args=[self.region.pk]),
            follow=True,
        )
        self.assertRedirects(blocked, reverse("region-list"))
        self.region.refresh_from_db()
        self.assertTrue(self.region.is_active)
        self.assertContains(blocked, "Aktif tank, tanker veya araç bağlıyken bölge pasife alınamaz")

        create_page = self.client.get(reverse("region-create"))
        self.assertContains(create_page, "Kısa ve ayırt edilebilir kod")

        response = self.client.post(
            reverse("region-create"),
            {"name": "Yeni Bölge", "code": "YB", "is_active": "on"},
        )
        self.assertRedirects(response, reverse("region-list"))
        self.assertTrue(Region.objects.filter(code="YB").exists())

    def test_supplier_list_uses_separate_create_page_and_sums_only_active_receipts(self):
        voided_receipt = create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.SUPPLIER_RECEIPT,
            data={
                "occurred_at": self.now,
                "supplier": self.supplier,
                "destination_storage": self.fixed,
                "liters": Decimal("50"),
                "purchase_unit_price": Decimal("42.0000"),
            },
        )
        voided_receipt.voided_at = self.now
        voided_receipt.voided_by = self.user
        voided_receipt.void_reason = "Test iptali"
        voided_receipt.save(update_fields=["voided_at", "voided_by", "void_reason"])
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("supplier-list"))
        supplier = next(item for item in list_page.context["items"] if item == self.supplier)

        self.assertEqual(list_page.status_code, 200)
        self.assertEqual(supplier.total_received, Decimal("1000"))
        self.assertContains(list_page, "supplier-list-panel")
        self.assertContains(list_page, "Toplam gelen")
        self.assertContains(list_page, "1.000 L")
        self.assertContains(list_page, "Tedarikçi adı ara")
        self.assertContains(list_page, f'href="{reverse("supplier-create")}"')
        self.assertNotContains(list_page, "Yeni tedarikçi</h2>")

        create_page = self.client.get(reverse("supplier-create"))
        self.assertContains(create_page, "Toplam otomatik hesaplanır")

        response = self.client.post(
            reverse("supplier-create"),
            {"name": "Yeni Tedarikçi", "is_active": "on"},
        )
        self.assertRedirects(response, reverse("supplier-list"))
        self.assertTrue(Supplier.objects.filter(name="Yeni Tedarikçi").exists())

    def test_stock_adjustment_form_uses_counted_stock_and_shows_current_stock(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("stock-adjustment-create"))
        form = response.context["form"]
        storage_status = response.context["storage_status"]

        self.assertEqual(form.fields["storage"].empty_label, "Tank veya tanker seçin")
        self.assertEqual(
            list(form.fields),
            ["occurred_at", "storage", "counted_stock", "note"],
        )
        self.assertEqual(form.fields["note"].widget.attrs["rows"], 3)
        self.assertEqual(storage_status[str(self.fixed.pk)]["current_stock"], "875.000")
        self.assertContains(response, 'id="operation-storage-status"')
        self.assertContains(response, 'data-storage-status="adjustment"')
        self.assertContains(response, "Sayılan stok (L)")

        invalid = self.client.post(
            reverse("stock-adjustment-create"),
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "storage": self.fixed.pk,
                "counted_stock": "875.000",
                "note": "Fiziksel sayım farkı",
            },
        )

        self.assertEqual(invalid.status_code, 200)
        self.assertFormError(
            invalid.context["form"],
            "counted_stock",
            "Sayılan stok mevcut stokla aynı; düzeltme gerekmiyor.",
        )
        self.assertFalse(
            FuelMovement.objects.filter(
                movement_type=FuelMovement.MovementType.STOCK_ADJUSTMENT
            ).exists()
        )

    def test_manager_can_create_stock_adjustment(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("stock-adjustment-create"),
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "storage": self.fixed.pk,
                "counted_stock": "1000.000",
                "note": "Fiziksel sayım fazlası",
            },
        )

        self.assertRedirects(response, reverse("movement-list"))
        adjustment = FuelMovement.objects.get(
            movement_type=FuelMovement.MovementType.STOCK_ADJUSTMENT
        )
        self.assertEqual(adjustment.destination_storage, self.fixed)
        self.assertEqual(adjustment.liters, Decimal("125"))
        self.assertIn("Sistem stoğu: 875.000 L", adjustment.note)
        self.assertIn("Sayılan stok: 1000.000 L", adjustment.note)
        dashboard = self.client.get(reverse("dashboard"))
        self.assertEqual(dashboard.context["total_stock"], Decimal("1000"))
        self.assertEqual(dashboard.context["totals"]["received"], Decimal("1000"))
        self.assertEqual(dashboard.context["totals"]["consumed"], Decimal("125"))

        movement_page = self.client.get(reverse("movement-list"))
        correction_url = reverse("movement-correct", args=[adjustment.pk])
        self.assertNotContains(movement_page, f'href="{correction_url}"')

        correction = self.client.get(correction_url, follow=True)
        self.assertRedirects(correction, reverse("movement-list"))
        self.assertContains(
            correction,
            "Stok düzeltmeleri fiziksel sayım kaydıdır. Hatalı kaydı iptal edip yeni sayım girin.",
        )

    def test_manager_can_void_movement_without_deleting_audit_record(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("movement-void", args=[self.fueling.pk]),
            {"reason": "Yanlış araç seçildi"},
        )

        self.assertRedirects(response, reverse("movement-list"))
        self.fueling.refresh_from_db()
        self.assertTrue(self.fueling.is_voided)
        self.assertEqual(self.fueling.voided_by, self.user)
        self.assertTrue(FuelMovement.objects.filter(pk=self.fueling.pk).exists())

        movement_page = self.client.get(reverse("movement-list"), {"status": "voided"})
        self.assertContains(movement_page, "Yanlış araç seçildi")
        dashboard = self.client.get(reverse("dashboard"))
        self.assertEqual(dashboard.context["total_stock"], Decimal("1000"))
        self.assertEqual(dashboard.context["totals"]["consumed"], Decimal("0"))

        export = self.client.get(reverse("report-export-xlsx"))
        workbook = load_workbook(BytesIO(export.content), data_only=True)
        summary_values = {
            row[0].value: row[1].value
            for row in workbook["Özet"].iter_rows(min_col=1, max_col=2)
            if row[0].value
        }
        self.assertEqual(summary_values["Araç tüketimi"], 0)
        self.assertEqual(workbook["Hareketler"].max_row, 2)

    def test_user_without_change_permission_cannot_void_movement(self):
        user = get_user_model().objects.create_user(
            username="viewer", password="test-password"
        )
        user.user_permissions.add(
            Permission.objects.get(
                content_type__app_label="fuel", codename="view_fuelmovement"
            )
        )
        self.client.force_login(user)

        response = self.client.get(reverse("movement-void", args=[self.fueling.pk]))

        self.assertEqual(response.status_code, 403)

    def test_excel_export_matches_report_totals(self):
        self.client.force_login(self.user)
        response = self.client.get(reverse("report-export-xlsx"))
        self.assertEqual(response.status_code, 200)
        workbook = load_workbook(BytesIO(response.content), data_only=True)
        summary = workbook["Özet"]
        summary_values = {
            row[0].value: row[1].value
            for row in summary.iter_rows(min_col=1, max_col=2)
            if row[0].value
        }
        self.assertEqual(summary_values["Tedarikçiden gelen"], 1000)
        self.assertEqual(summary_values["Araç tüketimi"], 125)
        self.assertEqual(summary_values["Dönem sonu stok"], 875)
        self.assertEqual(summary_values["Stok mutabakatı"], 0)
        self.assertEqual(workbook["Hareketler"].max_row, 3)
        self.assertIn("Bölge Özeti", workbook.sheetnames)
        self.assertIn("Stok Noktaları", workbook.sheetnames)
        self.assertIn("Tedarikçi Özeti", workbook.sheetnames)
        self.assertIn("Günlük Özet", workbook.sheetnames)
        self.assertEqual(workbook["Bölge Özeti"]["B2"].value, 1000)
        self.assertEqual(workbook["Bölge Özeti"]["E2"].value, 125)
        storage_sheet = workbook["Stok Noktaları"]
        storage_values = {
            cell.value: storage_sheet.cell(row=2, column=cell.column).value
            for cell in storage_sheet[1]
        }
        self.assertEqual(storage_values["Dönem başı stok"], 0)
        self.assertEqual(storage_values["Dönem sonu stok"], 875)
        self.assertEqual(storage_values["Anlık stok"], 875)
        self.assertEqual(len(summary._charts), 2)
        self.assertEqual(summary["B12"].number_format, '#,##0.###" L"')
        self.assertEqual(len(summary._charts[0].series), 1)
        self.assertEqual(len(summary._charts[1].series), 1)
        self.assertIsNone(workbook["Hareketler"].auto_filter.ref)
        self.assertEqual(len(workbook["Hareketler"].tables), 1)

    def test_excel_export_includes_period_stock_rfid_meter_and_scale_details(self):
        opening_at = self.now - timedelta(days=2)
        create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.OPENING_BALANCE,
            data={
                "occurred_at": opening_at,
                "destination_storage": self.fixed,
                "liters": Decimal("500"),
            },
        )
        self.receipt.external_scale_kg = Decimal("850")
        self.receipt.company_scale_kg = Decimal("847.5")
        self.receipt.save(update_fields=["external_scale_kg", "company_scale_kg"])
        self.vehicle.tag_code = "4202521459"
        self.vehicle.meter_type = Vehicle.MeterType.KM
        self.vehicle.save(update_fields=["tag_code", "meter_type"])
        self.fueling.meter_value = Decimal("1000")
        self.fueling.save(update_fields=["meter_value"])
        create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.VEHICLE_FUELING,
            data={
                # Kurulum dolumuyla aynı anda (daha büyük pk): sayaç
                # kronolojisi bozulmadan "son sayaç" kaydı olur.
                "occurred_at": self.now,
                "source_storage": self.fixed,
                "vehicle": self.vehicle,
                "liters": Decimal("50"),
                "meter_value": Decimal("1100"),
            },
        )
        start = (self.now - timedelta(days=1)).date().isoformat()
        end = self.now.date().isoformat()
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("report-export-xlsx"), {"start": start, "end": end}
        )

        self.assertEqual(response.status_code, 200)
        self.assertIn(
            f'akaryakit-raporu_{start.replace("-", "")}_{end.replace("-", "")}.xlsx',
            response["Content-Disposition"],
        )
        workbook = load_workbook(BytesIO(response.content), data_only=True)

        storage_sheet = workbook["Stok Noktaları"]
        storage_values = {
            cell.value: storage_sheet.cell(row=2, column=cell.column).value
            for cell in storage_sheet[1]
        }
        self.assertEqual(storage_values["Dönem başı stok"], 500)
        self.assertEqual(storage_values["Dönem sonu stok"], 1325)
        self.assertEqual(storage_values["Anlık stok"], 1325)

        vehicle_sheet = workbook["Araç Özeti"]
        vehicle_values = {
            cell.value: vehicle_sheet.cell(row=2, column=cell.column).value
            for cell in vehicle_sheet[1]
        }
        self.assertEqual(vehicle_values["RFID kodu"], "4202521459")
        self.assertEqual(vehicle_values["İlk sayaç"], 1000)
        self.assertEqual(vehicle_values["Son sayaç"], 1100)
        self.assertEqual(vehicle_values["Sayaç farkı"], 100)

        supplier_sheet = workbook["Tedarikçi Özeti"]
        supplier_values = {
            cell.value: supplier_sheet.cell(row=2, column=cell.column).value
            for cell in supplier_sheet[1]
        }
        self.assertEqual(supplier_values["Tedarikçi kantarı"], 850)
        self.assertEqual(supplier_values["Kurum kantarı"], 847.5)
        self.assertEqual(supplier_values["Kantar farkı"], 2.5)

        movement_headers = [cell.value for cell in workbook["Hareketler"][1]]
        self.assertIn("Araç RFID", movement_headers)
        self.assertIn("Tedarikçi kantarı", movement_headers)
        self.assertIn("Kurum kantarı", movement_headers)

    def test_reports_include_region_and_storage_summaries(self):
        self.client.force_login(self.user)

        response = self.client.get(reverse("reports"))

        self.assertEqual(response.status_code, 200)
        region_row = response.context["region_summary"][0]
        self.assertEqual(region_row["region_name"], self.region.name)
        self.assertEqual(region_row["received"], Decimal("1000"))
        self.assertEqual(region_row["consumed"], Decimal("125"))
        storage_row = response.context["storage_summary"][0]
        self.assertEqual(storage_row["storage"], self.fixed)
        self.assertEqual(storage_row["received"], Decimal("1000"))
        self.assertEqual(storage_row["vehicle_out"], Decimal("125"))
        self.assertEqual(storage_row["current_stock"], Decimal("875"))

    def test_user_without_view_permission_receives_403(self):
        user = get_user_model().objects.create_user(
            username="unauthorized", password="test-password"
        )
        self.client.force_login(user)
        response = self.client.get(reverse("dashboard"))
        self.assertEqual(response.status_code, 403)

    def test_manager_can_correct_movement_from_audited_form(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("movement-correct", args=[self.fueling.pk]),
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "source_storage": self.fixed.pk,
                "vehicle": self.vehicle.pk,
                "liters": "100.000",
                "meter_value": "",
                "note": "Düzeltilmiş miktar",
                "correction_reason": "Litre yanlış girildi",
            },
        )

        self.assertRedirects(response, reverse("movement-list"))
        self.fueling.refresh_from_db()
        self.assertTrue(self.fueling.is_voided)
        replacement = self.fueling.correction
        self.assertEqual(replacement.liters, Decimal("100"))
        self.assertEqual(replacement.corrects, self.fueling)

        movement_page = self.client.get(reverse("movement-list"))
        self.assertContains(movement_page, "Düzeltildi")
        self.assertContains(movement_page, "Düzelt")

    def test_vehicle_definition_can_be_edited_and_deactivated(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("vehicle-edit", args=[self.vehicle.pk]),
            {
                "name": "Güncel araç adı",
                "code": self.vehicle.code,
                "region": self.region.pk,
                "meter_type": self.vehicle.meter_type,
                "organization": "Saha firması",
                "is_active": "on",
            },
        )
        self.assertRedirects(response, reverse("vehicle-list"))
        self.vehicle.refresh_from_db()
        self.assertEqual(self.vehicle.name, "Güncel araç adı")

        response = self.client.post(reverse("vehicle-toggle", args=[self.vehicle.pk]))
        self.assertRedirects(response, reverse("vehicle-list"))
        self.vehicle.refresh_from_db()
        self.assertFalse(self.vehicle.is_active)

        fueling_page = self.client.get(reverse("vehicle-fueling-create"))
        self.assertNotContains(fueling_page, self.vehicle.code)

    def test_vehicle_list_uses_full_width_table_and_separate_create_page(self):
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("vehicle-list"))
        create_url = reverse("vehicle-create")
        self.assertContains(list_page, "vehicle-list-panel")
        self.assertContains(list_page, f'href="{create_url}"')
        self.assertContains(list_page, "Yeni araç veya makine")
        self.assertContains(list_page, "row-action-menu")
        self.assertNotContains(list_page, "Yeni araç/makine")

        create_page = self.client.get(create_url)
        form = create_page.context["form"]
        self.assertEqual(form.fields["region"].empty_label, "Bölge seçin")
        self.assertContains(create_page, "Sınıf ve sayaç ayrı bilgilerdir")

        response = self.client.post(
            create_url,
            {
                "name": "Yeni Kamyon",
                "code": "06YENI01",
                "region": self.region.pk,
                "meter_type": Vehicle.MeterType.KM,
                "organization": "Test Taşeron",
                "is_active": "on",
            },
        )
        self.assertRedirects(response, reverse("vehicle-list"))
        created = Vehicle.objects.get(code="06YENI01")
        self.assertEqual(created.name, "Yeni Kamyon")
        self.assertEqual(created.region, self.region)

    def test_vehicle_region_change_is_audited_and_preserves_historical_reports(self):
        new_region = Region.objects.create(name="2. Bölge", code="B2")
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("vehicle-list"))
        change_url = reverse("vehicle-region-change", args=[self.vehicle.pk])
        self.assertContains(list_page, f'href="{change_url}"')

        form_page = self.client.get(change_url)
        self.assertEqual(form_page.status_code, 200)
        self.assertContains(form_page, "Yakıt kaydı oluşturmaz")
        self.assertNotIn(
            self.region,
            form_page.context["form"].fields["to_region"].queryset,
        )

        response = self.client.post(
            change_url,
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "to_region": new_region.pk,
                "reason": "Araç yeni bölgeye görevlendirildi",
            },
        )
        self.assertRedirects(response, reverse("vehicle-list"))

        self.vehicle.refresh_from_db()
        self.fueling.refresh_from_db()
        change = VehicleRegionChange.objects.get(vehicle=self.vehicle)
        self.assertEqual(self.vehicle.region, new_region)
        self.assertEqual(change.from_region, self.region)
        self.assertEqual(change.to_region, new_region)
        self.assertEqual(change.created_by, self.user)
        self.assertEqual(self.fueling.vehicle_region, self.region)

        create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.VEHICLE_FUELING,
            data={
                "occurred_at": timezone.now(),
                "source_storage": self.fixed,
                "vehicle": self.vehicle,
                "liters": Decimal("50"),
            },
        )
        old_region_report = self.client.get(
            reverse("reports"),
            {"region": self.region.pk},
        )
        new_region_report = self.client.get(
            reverse("reports"),
            {"region": new_region.pk},
        )
        self.assertEqual(
            old_region_report.context["totals"]["consumed"],
            Decimal("125"),
        )
        self.assertEqual(
            new_region_report.context["totals"]["consumed"],
            Decimal("50"),
        )
        self.assertEqual(
            old_region_report.context["vehicle_summary"][0]["vehicle_region__name"],
            self.region.name,
        )
        self.assertEqual(
            new_region_report.context["vehicle_summary"][0]["vehicle_region__name"],
            new_region.name,
        )

        history_page = self.client.get(change_url)
        self.assertContains(history_page, "Araç yeni bölgeye görevlendirildi")
        self.assertContains(history_page, new_region.name)

    def test_storage_with_stock_cannot_be_deactivated(self):
        self.client.force_login(self.user)
        response = self.client.post(
            reverse("storage-toggle", args=[self.fixed.pk]),
            follow=True,
        )

        self.assertRedirects(response, reverse("storage-list"))
        self.fixed.refresh_from_db()
        self.assertTrue(self.fixed.is_active)
        self.assertContains(response, "Stoğu sıfır olmayan birim pasife alınamaz")

    def test_storage_list_uses_full_width_table_and_separate_create_page(self):
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("storage-list"))
        create_url = reverse("storage-create")
        self.assertContains(list_page, "storage-list-panel")
        self.assertContains(list_page, f'href="{create_url}"')
        self.assertContains(list_page, "Yeni tank veya tanker")
        self.assertContains(list_page, "row-action-menu")
        self.assertNotContains(list_page, "Yeni depolama birimi")

        create_page = self.client.get(create_url)
        form = create_page.context["form"]
        self.assertEqual(list(form.fields["kind"].choices)[0], ("", "Tür seçin"))
        self.assertEqual(form.fields["region"].empty_label, "Bölge seçin")
        self.assertContains(create_page, "Stok daha sonra girilir")

        response = self.client.post(
            create_url,
            {
                "name": "Yeni Sabit Tank",
                "code": "YENI-1",
                "kind": StorageUnit.Kind.FIXED_TANK,
                "region": self.region.pk,
                "capacity_liters": "5000.000",
                "low_stock_threshold": "250.000",
                "is_active": "on",
            },
        )
        self.assertRedirects(response, reverse("storage-list"))
        created = StorageUnit.objects.get(code="YENI-1")
        self.assertEqual(created.name, "Yeni Sabit Tank")
        self.assertEqual(created.region, self.region)

    def test_storage_list_filters_by_region(self):
        other_region = Region.objects.create(name="2. Bölge", code="B2")
        other_storage = StorageUnit.objects.create(
            name="İkinci Bölge Tankı",
            code="SABIT-2",
            kind=StorageUnit.Kind.FIXED_TANK,
            region=other_region,
            capacity_liters=Decimal("5000"),
        )
        self.client.force_login(self.user)

        response = self.client.get(
            reverse("storage-list"), {"region": other_region.pk}
        )

        self.assertEqual(response.context["region_filter"], str(other_region.pk))
        self.assertContains(response, other_storage.code)
        self.assertNotContains(response, self.fixed.code)

    def test_storage_region_change_is_audited_and_preserves_historical_reports(self):
        new_region = Region.objects.create(name="2. Bölge", code="B2")
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("storage-list"))
        change_url = reverse("storage-region-change", args=[self.fixed.pk])
        self.assertContains(list_page, f'href="{change_url}"')

        form_page = self.client.get(change_url)
        self.assertEqual(form_page.status_code, 200)
        self.assertContains(form_page, "Stok hareketi oluşturmaz")
        self.assertContains(form_page, self.region.name)
        self.assertNotIn(
            self.region,
            form_page.context["form"].fields["to_region"].queryset,
        )

        response = self.client.post(
            change_url,
            {
                "occurred_at": timezone.localtime(self.now).strftime("%Y-%m-%dT%H:%M"),
                "to_region": new_region.pk,
                "reason": "Eski bölge kapatıldı",
            },
        )

        self.assertRedirects(response, reverse("storage-list"))
        self.fixed.refresh_from_db()
        change = StorageRegionChange.objects.get(storage=self.fixed)
        self.assertEqual(self.fixed.region, new_region)
        self.assertEqual(change.from_region, self.region)
        self.assertEqual(change.to_region, new_region)
        self.assertEqual(change.created_by, self.user)
        self.assertEqual(change.reason, "Eski bölge kapatıldı")
        self.assertEqual(self.receipt.destination_region, self.region)
        self.assertEqual(self.fueling.source_region, self.region)

        create_movement(
            user=self.user,
            movement_type=FuelMovement.MovementType.SUPPLIER_RECEIPT,
            data={
                "occurred_at": timezone.now(),
                "supplier": self.supplier,
                "destination_storage": self.fixed,
                "liters": Decimal("100"),
                "purchase_unit_price": Decimal("42.0000"),
            },
        )
        old_region_report = self.client.get(
            reverse("reports"),
            {"region": self.region.pk},
        )
        new_region_report = self.client.get(
            reverse("reports"),
            {"region": new_region.pk},
        )
        self.assertEqual(
            old_region_report.context["totals"]["received"],
            Decimal("1000"),
        )
        self.assertEqual(
            new_region_report.context["totals"]["received"],
            Decimal("100"),
        )

        history_page = self.client.get(change_url)
        self.assertContains(history_page, "Eski bölge kapatıldı")
        self.assertContains(history_page, new_region.name)

    def test_vehicle_list_searches_code_name_and_organization(self):
        self.vehicle.organization = "Aranan Taşeron"
        self.vehicle.save(update_fields=["organization"])
        self.client.force_login(self.user)

        response = self.client.get(reverse("vehicle-list"), {"q": "Aranan"})

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, self.vehicle.code)

    def test_superuser_can_create_user_with_role(self):
        group = Group.objects.create(name="Veri Girişi")
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("user-create"),
            {
                "username": "operator2",
                "first_name": "Saha",
                "last_name": "Operatörü",
                "email": "operator@example.test",
                "groups": [group.pk],
                "is_active": "on",
                "password1": "GuvenliTest2026!",
                "password2": "GuvenliTest2026!",
            },
        )

        self.assertRedirects(response, reverse("user-list"))
        created = get_user_model().objects.get(username="operator2")
        self.assertTrue(created.check_password("GuvenliTest2026!"))
        self.assertTrue(created.groups.filter(pk=group.pk).exists())

    def test_user_list_is_full_width_and_create_form_uses_role_checkboxes(self):
        group = Group.objects.create(name="Rapor Görüntüleme")
        self.user.groups.add(group)
        self.client.force_login(self.user)

        list_page = self.client.get(reverse("user-list"))

        self.assertEqual(list_page.status_code, 200)
        self.assertContains(list_page, "user-list-panel")
        self.assertContains(list_page, "row-action-menu")
        self.assertContains(list_page, f'href="{reverse("user-create")}"')
        self.assertNotContains(list_page, "Yeni kullanıcı</h2>")

        create_page = self.client.get(reverse("user-create"))
        form = create_page.context["form"]

        self.assertEqual(form.fields["groups"].widget.__class__.__name__, "CheckboxSelectMultiple")
        self.assertEqual(form.fields["is_active"].widget.attrs["class"], "form-check-input")
        self.assertContains(create_page, "role-option-list")
        self.assertContains(create_page, "Kimlik bilgileri")
        self.assertContains(create_page, "Yetki ve hesap durumu")
        self.assertContains(create_page, "İlk parola")

    def test_non_superuser_cannot_open_user_management(self):
        operator = get_user_model().objects.create_user(
            username="operator3",
            password="test-password",
        )
        self.client.force_login(operator)

        response = self.client.get(reverse("user-list"))

        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.client.get(reverse("user-create")).status_code, 403)

    def test_superuser_cannot_deactivate_own_account(self):
        self.client.force_login(self.user)

        response = self.client.post(
            reverse("user-edit", args=[self.user.pk]),
            {
                "username": self.user.username,
                "first_name": "",
                "last_name": "",
                "email": self.user.email,
                "groups": [],
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertContains(response, "Kendi kullanıcı hesabınızı pasife alamazsınız")
        self.user.refresh_from_db()
        self.assertTrue(self.user.is_active)
