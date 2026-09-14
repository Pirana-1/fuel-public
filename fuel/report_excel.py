from datetime import timedelta
from decimal import Decimal

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import CharacterProperties, Paragraph, ParagraphProperties
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import FuelMovement, Vehicle
from .services import contractor_pricing_totals


ZERO = Decimal("0")
NAVY = "17324D"
NAVY_DARK = "10283F"
ORANGE = "F47A21"
BLUE_SOFT = "EAF2F8"
GREEN_SOFT = "E4F5EE"
RED_SOFT = "FDEAEA"
TEXT = "172434"
MUTED = "607286"
LINE = "D7E1E9"
LITERS_FORMAT = '#,##0.###" L"'
KG_FORMAT = '#,##0.###" kg"'
METER_FORMAT = "#,##0.##"
COUNT_FORMAT = "#,##0"
UNIT_PRICE_FORMAT = '#,##0.0000" TL/L"'
MONEY_FORMAT = '#,##0.00" TL"'


def _chart_axis_text():
    properties = CharacterProperties(solidFill=TEXT, sz=900)
    return RichText(
        p=[
            Paragraph(
                pPr=ParagraphProperties(defRPr=properties),
                endParaRPr=properties,
            )
        ]
    )


def report_filename(filters):
    start = (filters.get("start") or "baslangic").replace("-", "")
    end = (filters.get("end") or "bugun").replace("-", "")
    return f"akaryakit-raporu_{start}_{end}.xlsx"


def _format_period_date(value):
    parts = (value or "").split("-")
    if len(parts) == 3 and all(part.isdigit() for part in parts):
        return ".".join(reversed(parts))
    return value or "Sınırsız"


def _period_text(filters):
    return (
        f"{_format_period_date(filters.get('start'))} – "
        f"{_format_period_date(filters.get('end'))}"
    )


def _local_naive(value):
    if value is None:
        return None
    if timezone.is_aware(value):
        return timezone.localtime(value).replace(tzinfo=None)
    return value


def _style_header(sheet):
    fill = PatternFill("solid", fgColor=NAVY)
    font = Font(color="FFFFFF", bold=True, size=10)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30


def _style_data_sheet(
    sheet,
    *,
    headers,
    rows,
    widths,
    table_name,
    number_formats=None,
    total_row=None,
    freeze_panes="A2",
    wrap_text=True,
    zoom_scale=90,
):
    rows = list(rows)
    sheet.append(headers)
    for row in rows:
        sheet.append(row)

    _style_header(sheet)
    sheet.freeze_panes = freeze_panes
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = zoom_scale
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape" if len(headers) > 6 else "portrait"
    sheet.print_title_rows = "1:1"
    sheet.oddHeader.center.text = f"&B{sheet.title}"
    sheet.oddFooter.center.text = "Sayfa &P / &N"

    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    last_data_row = len(rows) + 1
    if rows:
        table = Table(
            displayName=table_name,
            ref=f"A1:{get_column_letter(len(headers))}{last_data_row}",
        )
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
        sheet.append(["Kayıt bulunamadı"] + [None] * (len(headers) - 1))
        sheet.merge_cells(
            start_row=2,
            start_column=1,
            end_row=2,
            end_column=len(headers),
        )
        sheet["A2"].font = Font(italic=True, color=MUTED)
        sheet["A2"].alignment = Alignment(horizontal="center")
        sheet.row_dimensions[2].height = 28

    total_row_number = None
    if total_row is not None:
        total_row_number = sheet.max_row + 1
        sheet.append(total_row)
        for cell in sheet[total_row_number]:
            cell.fill = PatternFill("solid", fgColor=BLUE_SOFT)
            cell.font = Font(bold=True, color=TEXT)
            cell.border = Border(top=Side(style="thin", color=NAVY))

    number_formats = number_formats or {}
    final_row = total_row_number or last_data_row
    for column, number_format in number_formats.items():
        for row_number in range(2, final_row + 1):
            cell = sheet.cell(row=row_number, column=column)
            if cell.value is not None:
                cell.number_format = number_format
                cell.alignment = Alignment(horizontal="right", vertical="center")

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for cell in row:
            if cell.alignment.horizontal is None:
                cell.alignment = Alignment(vertical="center", wrap_text=wrap_text)
    if not wrap_text:
        for row_number in range(2, sheet.max_row + 1):
            sheet.row_dimensions[row_number].height = 22

    sheet.print_area = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"
    return len(rows)


def _movement_rows(movements):
    rows = []
    for movement in movements:
        meter_unit = ""
        if movement.vehicle_id:
            if movement.vehicle.meter_type == Vehicle.MeterType.KM:
                meter_unit = "km"
            elif movement.vehicle.meter_type == Vehicle.MeterType.HOUR:
                meter_unit = "saat"
        target_region = movement.destination_region or movement.vehicle_region
        operation_region = target_region or movement.source_region
        source_label = str(movement.source_storage or movement.supplier or "")
        target_label = str(movement.vehicle or movement.destination_storage or "")
        scale_difference = (
            movement.external_scale_kg - movement.company_scale_kg
            if movement.external_scale_kg is not None
            and movement.company_scale_kg is not None
            else None
        )
        rows.append(
            [
                _local_naive(movement.occurred_at),
                movement.get_movement_type_display(),
                source_label,
                target_label,
                movement.liters,
                movement.meter_value,
                meter_unit,
                operation_region.name if operation_region else "",
                movement.note,
                str(movement.supplier or ""),
                movement.delivery_note,
                movement.vehicle.tag_code if movement.vehicle_id else "",
                movement.created_by.get_username(),
                movement.source_region.name if movement.source_region_id else "",
                target_region.name if target_region else "",
                movement.external_scale_kg,
                movement.company_scale_kg,
                scale_difference,
                movement.pk,
                _local_naive(movement.created_at),
                (
                    movement.vehicle.get_meter_type_display()
                    if movement.vehicle_id
                    else ""
                ),
                (
                    "Taşeron"
                    if movement.vehicle_id and movement.vehicle.is_contractor
                    else "Kurum aracı" if movement.vehicle_id else ""
                ),
                movement.purchase_unit_price,
                movement.purchase_total,
                movement.contractor_reference_unit_price,
                movement.contractor_unit_price,
                movement.contractor_unit_difference,
                movement.contractor_reference_total,
                movement.contractor_charged_total,
                movement.contractor_difference_total,
            ]
        )
    return rows


def _supplier_receipt_rows(movements):
    rows = []
    for movement in movements:
        if movement.movement_type != FuelMovement.MovementType.SUPPLIER_RECEIPT:
            continue
        scale_difference = (
            movement.external_scale_kg - movement.company_scale_kg
            if movement.external_scale_kg is not None
            and movement.company_scale_kg is not None
            else None
        )
        rows.append(
            [
                _local_naive(movement.occurred_at),
                str(movement.supplier or ""),
                str(movement.destination_storage or ""),
                movement.destination_region.name if movement.destination_region_id else "",
                movement.liters,
                movement.purchase_unit_price,
                movement.purchase_total,
                movement.delivery_note,
                movement.external_scale_kg,
                movement.company_scale_kg,
                scale_difference,
                movement.created_by.get_username(),
                movement.note,
            ]
        )
    return rows


def _contractor_rows(movements):
    return [
        [
            _local_naive(movement.occurred_at),
            movement.vehicle.code,
            movement.vehicle.name,
            movement.vehicle.organization,
            movement.vehicle_region.name if movement.vehicle_region_id else "",
            str(movement.source_storage or ""),
            movement.liters,
            movement.contractor_reference_unit_price,
            movement.contractor_unit_price,
            movement.contractor_unit_difference,
            movement.contractor_reference_total,
            movement.contractor_charged_total,
            movement.contractor_difference_total,
            movement.created_by.get_username(),
            movement.note,
        ]
        for movement in movements
    ]


def _build_contractor_summary_sheet(sheet, *, context, generated_by, generated_at):
    filters = context["filters"]
    financials = contractor_pricing_totals(context["queryset"])
    sheet.title = "Özet"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = ORANGE
    sheet.merge_cells("A1:F1")
    sheet["A1"] = "TAŞERON YAKIT FİYAT RAPORU"
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY_DARK)
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:F2")
    sheet["A2"] = f"Rapor dönemi: {_period_text(filters)}"
    sheet["A2"].font = Font(color=MUTED, italic=True, size=10)

    sheet["A4"] = "FİNANSAL ÖZET"
    sheet["A4"].font = Font(bold=True, color=ORANGE)
    metrics = [
        ("Toplam verilen yakıt", sum((item.liters for item in context["queryset"]), ZERO), LITERS_FORMAT),
        ("Referans maliyet", financials["reference_total"], MONEY_FORMAT),
        ("Yansıtılan tutar", financials["charged_total"], MONEY_FORMAT),
        ("Toplam fark", financials["difference_total"], MONEY_FORMAT),
        ("Ağırlıklı ortalama taşeron fiyatı", financials["weighted_average_unit_price"], UNIT_PRICE_FORMAT),
        ("Fiyatlı dolum sayısı", financials["priced_fueling_count"], COUNT_FORMAT),
        ("Fiyatsız eski dolum sayısı", financials["unpriced_fueling_count"], COUNT_FORMAT),
    ]
    for row_number, (label, value, number_format) in enumerate(metrics, start=5):
        label_cell = sheet.cell(row=row_number, column=1, value=label)
        value_cell = sheet.cell(row=row_number, column=2, value=value)
        label_cell.fill = PatternFill("solid", fgColor=BLUE_SOFT)
        label_cell.font = Font(bold=True, color=TEXT)
        value_cell.fill = PatternFill("solid", fgColor="F7FAFC")
        value_cell.font = Font(bold=True, color=NAVY)
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="right")

    sheet["A14"] = "RAPOR BİLGİSİ"
    sheet["A14"].font = Font(bold=True, color=ORANGE)
    info_rows = [
        ("Hazırlayan", generated_by.get_full_name().strip() or generated_by.get_username()),
        ("Oluşturulma", timezone.localtime(generated_at).strftime("%d.%m.%Y %H:%M")),
    ]
    for row_number, (label, value) in enumerate(info_rows, start=15):
        sheet.cell(row=row_number, column=1, value=label).font = Font(bold=True, color=TEXT)
        sheet.cell(row=row_number, column=2, value=value).font = Font(color=MUTED)
    sheet.merge_cells("A18:F20")
    sheet["A18"] = (
        "Fiyatlar her dolum kaydına işlem anında sabitlenir. Sonradan yeni bir fiyat "
        "tanımlanması geçmiş dolumların tutarını değiştirmez."
    )
    sheet["A18"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A18"].font = Font(color=MUTED, size=9)
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 28
    for column in range(3, 7):
        sheet.column_dimensions[get_column_letter(column)].width = 14
    sheet.freeze_panes = "A4"
    sheet.print_area = "A1:F20"


def _vehicle_meter_bounds(movements):
    bounds = {}
    ordered = sorted(movements, key=lambda item: (item.occurred_at, item.pk))
    for movement in ordered:
        if (
            movement.movement_type != FuelMovement.MovementType.VEHICLE_FUELING
            or not movement.vehicle_id
            or movement.meter_value is None
        ):
            continue
        key = (movement.vehicle_id, movement.vehicle_region_id)
        current = bounds.setdefault(
            key,
            {"first": movement.meter_value, "last": movement.meter_value},
        )
        current["last"] = movement.meter_value
    return bounds


def _vehicle_rows(summary, movements):
    meter_labels = dict(Vehicle.MeterType.choices)
    bounds = _vehicle_meter_bounds(movements)
    rows = []
    for item in summary:
        meter = bounds.get((item["vehicle_id"], item["vehicle_region_id"]), {})
        first_meter = meter.get("first")
        last_meter = meter.get("last")
        meter_difference = (
            last_meter - first_meter
            if first_meter is not None
            and last_meter is not None
            and last_meter >= first_meter
            else None
        )
        rows.append(
            [
                item["vehicle__code"],
                item["vehicle__name"],
                item["vehicle__tag_code"],
                item["vehicle_region__name"],
                item["vehicle__organization"],
                "Taşeron" if item["vehicle__is_contractor"] else "Kurum aracı",
                meter_labels.get(item["vehicle__meter_type"], ""),
                item["fueling_count"],
                item["total_liters"],
                first_meter,
                last_meter,
                meter_difference,
            ]
        )
    return rows


def _sum(rows, key):
    return sum((row.get(key) or ZERO for row in rows), ZERO)


def _weekly_consumption_rows(daily_summary):
    daily_summary = list(daily_summary)
    if not daily_summary:
        return []
    first_day = daily_summary[0]["day"]
    last_day = daily_summary[-1]["day"]
    totals = {}
    for item in daily_summary:
        bucket = (item["day"] - first_day).days // 7
        totals[bucket] = totals.get(bucket, ZERO) + (item["consumed"] or ZERO)
    rows = []
    for bucket, total in sorted(totals.items()):
        bucket_start = first_day + timedelta(days=bucket * 7)
        bucket_end = min(bucket_start + timedelta(days=6), last_day)
        rows.append(
            [
                f"{bucket_start:%d.%m}–{bucket_end:%d.%m}",
                total,
            ]
        )
    return rows


def _build_summary_sheet(
    sheet,
    *,
    context,
    generated_by,
    generated_at,
    region_sheet,
    region_row_count,
    daily_sheet,
    weekly_row_count,
):
    if context.get("contractor_report"):
        _build_contractor_summary_sheet(
            sheet,
            context=context,
            generated_by=generated_by,
            generated_at=generated_at,
        )
        return
    filters = context["filters"]
    labels = context["filter_labels"]
    storage_rows = context["storage_summary"]
    totals = context["totals"]

    opening_stock = _sum(storage_rows, "opening_stock")
    opening_added = _sum(storage_rows, "opening_added")
    adjustment_net = _sum(storage_rows, "adjustment_net")
    closing_stock = _sum(storage_rows, "closing_stock")
    current_stock = _sum(storage_rows, "current_stock")
    transfer_in = _sum(storage_rows, "transfer_in")
    transfer_out = _sum(storage_rows, "transfer_out")
    vehicle_out = _sum(storage_rows, "vehicle_out")
    storage_received = _sum(storage_rows, "received")

    can_reconcile = not filters["region"] and not filters["vehicle"]
    expected_closing = (
        opening_stock
        + opening_added
        + storage_received
        + transfer_in
        - transfer_out
        - vehicle_out
        + adjustment_net
    )
    reconciliation_difference = closing_stock - expected_closing

    sheet.title = "Özet"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 90
    sheet.sheet_properties.tabColor = ORANGE
    sheet.merge_cells("A1:K1")
    sheet["A1"] = "AKARYAKIT HAREKET VE STOK RAPORU"
    sheet["A1"].fill = PatternFill("solid", fgColor=NAVY_DARK)
    sheet["A1"].font = Font(color="FFFFFF", bold=True, size=18)
    sheet["A1"].alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 34
    sheet.merge_cells("A2:K2")
    sheet["A2"] = f"Rapor dönemi: {_period_text(filters)}"
    sheet["A2"].font = Font(color=MUTED, italic=True, size=10)
    sheet["A2"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[2].height = 24

    info_rows = [
        ("Bölge", labels["region"]),
        ("Tank / tanker", labels["storage"]),
        ("Araç / makine", labels["vehicle"]),
        ("Hazırlayan", generated_by.get_full_name().strip() or generated_by.get_username()),
        ("Oluşturulma", timezone.localtime(generated_at).strftime("%d.%m.%Y %H:%M")),
    ]
    if filters.get("meter_type"):
        info_rows.insert(3, ("Sayaç türü", labels["meter_type"]))
    if filters.get("vehicle_class"):
        info_rows.insert(3, ("Araç sınıfı", labels["vehicle_class"]))
    sheet["A4"] = "RAPOR KAPSAMI"
    sheet["A4"].font = Font(bold=True, color=ORANGE)
    for row_number, (label, value) in enumerate(info_rows, start=5):
        sheet.cell(row=row_number, column=1, value=label).font = Font(
            bold=True, color=TEXT
        )
        sheet.cell(row=row_number, column=2, value=value).font = Font(color=MUTED)

    sheet["A11"] = "YÖNETİCİ ÖZETİ"
    sheet["A11"].font = Font(bold=True, color=ORANGE)
    metrics = [
        ("Dönem başı stok", opening_stock, LITERS_FORMAT),
        ("Dönem içi başlangıç stoğu", opening_added, LITERS_FORMAT),
        ("Tedarikçiden gelen", totals["received"], LITERS_FORMAT),
        ("Toplam alış tutarı", totals["purchase_amount"], MONEY_FORMAT),
        ("Araç tüketimi", totals["consumed"], LITERS_FORMAT),
        ("Stok düzeltmesi net", adjustment_net, LITERS_FORMAT),
        ("Dönem sonu stok", closing_stock, LITERS_FORMAT),
        ("Anlık stok", current_stock, LITERS_FORMAT),
        ("İç transfer", totals["transferred"], LITERS_FORMAT),
    ]
    for row_number, (label, value, number_format) in enumerate(metrics, start=12):
        label_cell = sheet.cell(row=row_number, column=1, value=label)
        value_cell = sheet.cell(row=row_number, column=2, value=value)
        label_cell.fill = PatternFill("solid", fgColor=BLUE_SOFT)
        label_cell.font = Font(bold=True, color=TEXT)
        value_cell.fill = PatternFill("solid", fgColor="F7FAFC")
        value_cell.font = Font(bold=True, color=NAVY)
        value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="right")

    reconciliation_row = 21
    sheet.cell(reconciliation_row, 1, "Stok mutabakatı")
    sheet.cell(reconciliation_row, 1).font = Font(bold=True, color=TEXT)
    if can_reconcile:
        cell = sheet.cell(reconciliation_row, 2, reconciliation_difference)
        cell.number_format = LITERS_FORMAT
        cell.font = Font(
            bold=True,
            color="087553" if reconciliation_difference == ZERO else "B42318",
        )
        cell.fill = PatternFill(
            "solid",
            fgColor=GREEN_SOFT if reconciliation_difference == ZERO else RED_SOFT,
        )
    else:
        sheet.cell(reconciliation_row, 2, "Bölge/araç filtresinde uygulanmaz")
        sheet.cell(reconciliation_row, 2).font = Font(italic=True, color=MUTED)

    sheet.merge_cells("A23:B25")
    sheet["A23"] = (
        "Dönem sonu stok seçilen dönemin son anını, anlık stok ise dosyanın "
        "oluşturulduğu andaki gerçek bakiyeyi gösterir. İç transfer kurum toplam "
        "stoğunu ve tüketimi değiştirmez."
    )
    sheet["A23"].alignment = Alignment(wrap_text=True, vertical="top")
    sheet["A23"].font = Font(color=MUTED, size=9)

    sheet.column_dimensions["A"].width = 29
    sheet.column_dimensions["B"].width = 27
    sheet.column_dimensions["C"].width = 3
    for column in range(4, 12):
        sheet.column_dimensions[get_column_letter(column)].width = 13

    if region_row_count:
        chart = BarChart()
        chart.type = "col"
        chart.grouping = "clustered"
        chart.style = 10
        chart.title = "Bölge Bazlı Araç Tüketimi (L)"
        chart.height = 7.2
        chart.width = 14.5
        chart.legend = None
        chart.add_data(
            Reference(
                region_sheet,
                min_col=5,
                min_row=1,
                max_row=region_row_count + 1,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                region_sheet,
                min_col=1,
                min_row=2,
                max_row=region_row_count + 1,
            )
        )
        chart.y_axis.numFmt = '#,##0" L"'
        chart.y_axis.majorGridlines = None
        chart.x_axis.txPr = _chart_axis_text()
        chart.y_axis.txPr = _chart_axis_text()
        chart.dLbls = DataLabelList()
        chart.dLbls.showVal = True
        chart.dLbls.showCatName = True
        chart.dLbls.showSerName = False
        chart.dLbls.showLegendKey = False
        chart.dLbls.separator = "\n"
        chart.dLbls.numFmt = '#,##0" L"'
        chart.series[0].graphicalProperties.solidFill = ORANGE
        chart.series[0].graphicalProperties.line.solidFill = ORANGE
        sheet.add_chart(chart, "D4")

    if weekly_row_count:
        chart = LineChart()
        chart.style = 10
        chart.title = "Haftalık Araç Tüketimi (L)"
        chart.height = 7.2
        chart.width = 14.5
        chart.add_data(
            Reference(
                daily_sheet,
                min_col=8,
                max_col=8,
                min_row=1,
                max_row=weekly_row_count + 1,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(
                daily_sheet,
                min_col=7,
                min_row=2,
                max_row=weekly_row_count + 1,
            )
        )
        chart.legend = None
        chart.y_axis.numFmt = '#,##0" L"'
        chart.y_axis.majorGridlines = None
        chart.y_axis.title = "Litre"
        chart.x_axis.title = "Tarih"
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.tickLblPos = "nextTo"
        chart.x_axis.txPr = _chart_axis_text()
        chart.y_axis.txPr = _chart_axis_text()
        chart.series[0].graphicalProperties.line.solidFill = ORANGE
        chart.series[0].graphicalProperties.line.width = 28575
        chart.series[0].marker.symbol = "circle"
        chart.series[0].marker.size = 5
        chart.series[0].marker.graphicalProperties.solidFill = ORANGE
        chart.series[0].marker.graphicalProperties.line.solidFill = ORANGE
        chart.series[0].smooth = True
        chart.dLbls = DataLabelList(
            dLblPos="b",
            showCatName=True,
            showLegendKey=False,
            showSerName=False,
            showVal=True,
            separator="\n",
            numFmt=LITERS_FORMAT,
        )
        sheet.add_chart(chart, "D20")

    sheet.freeze_panes = "A4"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.orientation = "landscape"
    sheet.print_area = "A1:K35"
    sheet.oddFooter.center.text = "Sayfa &P / &N"


def build_report_workbook(*, context, generated_by, generated_at):
    movements = context["queryset"]
    workbook = Workbook()
    summary_sheet = workbook.active
    workbook.properties.creator = generated_by.get_username()
    workbook.properties.title = "Akaryakıt hareket ve stok raporu"
    workbook.properties.subject = _period_text(context["filters"])
    workbook.properties.description = (
        "Yakıt hareketleri, stok noktaları, araçlar, tedarikçiler ve bölge özetleri"
    )

    movement_headers = [
        "Tarih",
        "İşlem",
        "Kaynak",
        "Hedef / araç",
        "Miktar",
        "Sayaç",
        "Sayaç birimi",
        "İşlem bölgesi",
        "Açıklama",
        "Tedarikçi",
        "İrsaliye",
        "Araç RFID",
        "Kaydı oluşturan",
        "Kaynak bölge",
        "Hedef / araç bölgesi",
        "Tedarikçi kantarı",
        "Kurum kantarı",
        "Kantar farkı",
        "Kayıt no",
        "Sisteme kayıt zamanı",
        "Sayaç türü",
        "Araç sınıfı",
        "Alış birim fiyatı",
        "Alış toplamı",
        "Referans ortalama fiyat",
        "Taşeron birim fiyatı",
        "Birim fiyat farkı",
        "Referans toplamı",
        "Yansıtılan toplam",
        "Toplam fiyat farkı",
    ]
    movement_sheet = workbook.create_sheet("Hareketler")
    movement_sheet.sheet_properties.tabColor = "4F81BD"
    _style_data_sheet(
        movement_sheet,
        headers=movement_headers,
        rows=_movement_rows(movements),
        widths=[18, 21, 27, 29, 14, 15, 12, 19, 34, 22, 18, 18, 18, 20, 23, 20, 18, 18, 12, 20, 18, 16, 20, 20, 23, 21, 19, 20, 21, 22],
        table_name="HareketlerTablosu",
        number_formats={
            1: "dd.mm.yyyy hh:mm",
            5: LITERS_FORMAT,
            6: METER_FORMAT,
            16: KG_FORMAT,
            17: KG_FORMAT,
            18: KG_FORMAT,
            19: COUNT_FORMAT,
            20: "dd.mm.yyyy hh:mm",
            23: UNIT_PRICE_FORMAT,
            24: MONEY_FORMAT,
            25: UNIT_PRICE_FORMAT,
            26: UNIT_PRICE_FORMAT,
            27: UNIT_PRICE_FORMAT,
            28: MONEY_FORMAT,
            29: MONEY_FORMAT,
            30: MONEY_FORMAT,
        },
        freeze_panes="C2",
        wrap_text=False,
        zoom_scale=85,
    )

    contractor_movements = [
        movement
        for movement in movements
        if movement.vehicle_id
        and movement.vehicle.is_contractor
        and movement.movement_type == FuelMovement.MovementType.VEHICLE_FUELING
    ]
    contractor_sheet = workbook.create_sheet("Taşeron Dolumları")
    contractor_sheet.sheet_properties.tabColor = ORANGE
    _style_data_sheet(
        contractor_sheet,
        headers=[
            "Tarih",
            "Araç kodu",
            "Araç / makine",
            "Bağlı firma / taşeron",
            "Bölge",
            "Kaynak tank / tanker",
            "Verilen yakıt",
            "Referans ortalama fiyat",
            "Taşeron birim fiyatı",
            "Birim fark",
            "Referans toplamı",
            "Yansıtılan toplam",
            "Toplam fark",
            "Kaydı oluşturan",
            "Açıklama",
        ],
        rows=_contractor_rows(contractor_movements),
        widths=[18, 18, 28, 26, 20, 28, 17, 23, 21, 18, 20, 21, 19, 18, 34],
        table_name="TaseronDolumlariTablosu",
        number_formats={
            1: "dd.mm.yyyy hh:mm",
            7: LITERS_FORMAT,
            8: UNIT_PRICE_FORMAT,
            9: UNIT_PRICE_FORMAT,
            10: UNIT_PRICE_FORMAT,
            11: MONEY_FORMAT,
            12: MONEY_FORMAT,
            13: MONEY_FORMAT,
        },
        total_row=[
            "GENEL TOPLAM",
            None,
            None,
            None,
            None,
            None,
            sum((item.liters for item in contractor_movements), ZERO),
            None,
            None,
            None,
            sum((item.contractor_reference_total or ZERO for item in contractor_movements), ZERO),
            sum((item.contractor_charged_total or ZERO for item in contractor_movements), ZERO),
            sum((item.contractor_difference_total or ZERO for item in contractor_movements), ZERO),
            None,
            None,
        ],
        freeze_panes="C2",
        wrap_text=False,
        zoom_scale=85,
    )

    vehicle_rows = _vehicle_rows(context["vehicle_summary"], movements)
    vehicle_sheet = workbook.create_sheet("Araç Özeti")
    vehicle_sheet.sheet_properties.tabColor = "F4A261"
    _style_data_sheet(
        vehicle_sheet,
        headers=[
            "Araç kodu",
            "Araç / makine",
            "RFID kodu",
            "Bölge",
            "Açıklama",
            "Araç sınıfı",
            "Sayaç türü",
            "Dolum sayısı",
            "Toplam litre",
            "İlk sayaç",
            "Son sayaç",
            "Sayaç farkı",
        ],
        rows=vehicle_rows,
        widths=[18, 28, 18, 22, 30, 16, 18, 15, 18, 16, 16, 16],
        table_name="AracOzetiTablosu",
        number_formats={
            8: COUNT_FORMAT,
            9: LITERS_FORMAT,
            10: METER_FORMAT,
            11: METER_FORMAT,
            12: METER_FORMAT,
        },
        total_row=[
            "GENEL TOPLAM",
            None,
            None,
            None,
            None,
            None,
            None,
            sum((row[7] for row in vehicle_rows), 0),
            sum((row[8] for row in vehicle_rows), ZERO),
            None,
            None,
            None,
        ],
    )

    supplier_rows = [
        [
            item["supplier__name"],
            item["delivery_count"],
            item["total_liters"],
            item["total_external_scale_kg"],
            item["total_company_scale_kg"],
            item["scale_pair_count"],
            item["scale_difference_kg"],
            item["average_unit_price"],
            item["total_purchase_amount"],
        ]
        for item in context["supplier_summary"]
    ]
    supplier_sheet = workbook.create_sheet("Tedarikçi Özeti")
    supplier_sheet.sheet_properties.tabColor = "2A9D8F"
    _style_data_sheet(
        supplier_sheet,
        headers=[
            "Tedarikçi",
            "Teslimat sayısı",
            "Toplam litre",
            "Tedarikçi kantarı",
            "Kurum kantarı",
            "İki kantarı da girilen",
            "Kantar farkı",
            "Ağırlıklı ortalama alış fiyatı",
            "Toplam alış tutarı",
        ],
        rows=supplier_rows,
        widths=[28, 18, 19, 21, 19, 22, 18, 29, 22],
        table_name="TedarikciOzetiTablosu",
        number_formats={
            2: COUNT_FORMAT,
            3: LITERS_FORMAT,
            4: KG_FORMAT,
            5: KG_FORMAT,
            6: COUNT_FORMAT,
            7: KG_FORMAT,
            8: UNIT_PRICE_FORMAT,
            9: MONEY_FORMAT,
        },
        total_row=[
            "GENEL TOPLAM",
            sum((row[1] for row in supplier_rows), 0),
            sum((row[2] or ZERO for row in supplier_rows), ZERO),
            sum((row[3] or ZERO for row in supplier_rows), ZERO),
            sum((row[4] or ZERO for row in supplier_rows), ZERO),
            sum((row[5] for row in supplier_rows), 0),
            sum((row[6] or ZERO for row in supplier_rows), ZERO),
            (
                sum((row[8] or ZERO for row in supplier_rows), ZERO)
                / sum(
                    (
                        (item["priced_liters"] or ZERO)
                        for item in context["supplier_summary"]
                    ),
                    ZERO,
                )
                if sum(
                    (
                        (item["priced_liters"] or ZERO)
                        for item in context["supplier_summary"]
                    ),
                    ZERO,
                )
                else None
            ),
            sum((row[8] or ZERO for row in supplier_rows), ZERO),
        ],
    )

    supplier_receipt_rows = _supplier_receipt_rows(movements)
    supplier_receipt_sheet = workbook.create_sheet("Tedarikçi Alımları")
    supplier_receipt_sheet.sheet_properties.tabColor = "58A68C"
    _style_data_sheet(
        supplier_receipt_sheet,
        headers=[
            "Tarih",
            "Tedarikçi",
            "Yakıtın girdiği tank",
            "Bölge",
            "Miktar",
            "Alış birim fiyatı",
            "Alış toplamı",
            "İrsaliye",
            "Tedarikçi kantarı",
            "Kurum kantarı",
            "Kantar farkı",
            "Kaydı oluşturan",
            "Açıklama",
        ],
        rows=supplier_receipt_rows,
        widths=[18, 26, 29, 20, 17, 21, 20, 18, 21, 19, 18, 18, 34],
        table_name="TedarikciAlimlariTablosu",
        number_formats={
            1: "dd.mm.yyyy hh:mm",
            5: LITERS_FORMAT,
            6: UNIT_PRICE_FORMAT,
            7: MONEY_FORMAT,
            9: KG_FORMAT,
            10: KG_FORMAT,
            11: KG_FORMAT,
        },
        total_row=[
            "GENEL TOPLAM",
            None,
            None,
            None,
            sum((row[4] or ZERO for row in supplier_receipt_rows), ZERO),
            None,
            sum((row[6] or ZERO for row in supplier_receipt_rows), ZERO),
            None,
            sum((row[8] or ZERO for row in supplier_receipt_rows), ZERO),
            sum((row[9] or ZERO for row in supplier_receipt_rows), ZERO),
            sum((row[10] or ZERO for row in supplier_receipt_rows), ZERO),
            None,
            None,
        ],
        freeze_panes="C2",
        wrap_text=False,
        zoom_scale=85,
    )

    if context.get("contractor_report"):
        price_sheet = workbook.create_sheet("Fiyat Tanımları")
        price_sheet.sheet_properties.tabColor = "D6A84B"
        _style_data_sheet(
            price_sheet,
            headers=[
                "Geçerlilik başlangıcı",
                "Referans ortalama fiyat",
                "Taşeron birim fiyatı",
                "Birim fark",
                "Kaydı oluşturan",
                "Kayıt zamanı",
                "Açıklama",
            ],
            rows=[
                [
                    _local_naive(item.effective_from),
                    item.reference_unit_price,
                    item.contractor_unit_price,
                    item.unit_difference,
                    item.created_by.get_username(),
                    _local_naive(item.created_at),
                    item.note,
                ]
                for item in context["contractor_prices"]
            ],
            widths=[22, 24, 22, 18, 19, 20, 36],
            table_name="FiyatTanimlariTablosu",
            number_formats={
                1: "dd.mm.yyyy hh:mm",
                2: UNIT_PRICE_FORMAT,
                3: UNIT_PRICE_FORMAT,
                4: UNIT_PRICE_FORMAT,
                6: "dd.mm.yyyy hh:mm",
            },
        )

    region_rows = sorted(
        [
        [
            item["region_name"],
            item["received"],
            item["transfer_in"],
            item["transfer_out"],
            item["consumed"],
        ]
        for item in context["region_summary"]
        ],
        key=lambda row: row[4],
        reverse=True,
    )
    region_sheet = workbook.create_sheet("Bölge Özeti")
    region_sheet.sheet_properties.tabColor = "7A9E3A"
    region_row_count = _style_data_sheet(
        region_sheet,
        headers=[
            "Bölge",
            "Tedarikçiden gelen",
            "Transfer giriş",
            "Transfer çıkış",
            "Araç tüketimi",
        ],
        rows=region_rows,
        widths=[28, 22, 19, 19, 20],
        table_name="BolgeOzetiTablosu",
        number_formats={2: LITERS_FORMAT, 3: LITERS_FORMAT, 4: LITERS_FORMAT, 5: LITERS_FORMAT},
        total_row=[
            "GENEL TOPLAM",
            sum((row[1] for row in region_rows), ZERO),
            sum((row[2] for row in region_rows), ZERO),
            sum((row[3] for row in region_rows), ZERO),
            sum((row[4] for row in region_rows), ZERO),
        ],
    )

    storage_rows = []
    for item in context["storage_summary"]:
        storage = item["storage"]
        storage_rows.append(
            [
                storage.code,
                storage.name,
                storage.get_kind_display(),
                item["period_region"].name,
                item["opening_stock"],
                item["opening_added"],
                item["received"],
                item["transfer_in"],
                item["transfer_out"],
                item["vehicle_out"],
                item["adjustment_net"],
                item["closing_stock"],
                item["current_stock"],
                "Eksi stok" if item["is_negative_stock"] else "Normal",
            ]
        )
    storage_sheet = workbook.create_sheet("Stok Noktaları")
    storage_sheet.sheet_properties.tabColor = "6C8EBF"
    _style_data_sheet(
        storage_sheet,
        headers=[
            "Kod",
            "Tank / tanker",
            "Tür",
            "Dönem sonu bölgesi",
            "Dönem başı stok",
            "Dönem içi başlangıç stoğu",
            "Tedarikçiden gelen",
            "Transfer giriş",
            "Transfer çıkış",
            "Araca verilen",
            "Düzeltme net",
            "Dönem sonu stok",
            "Anlık stok",
            "Stok durumu",
        ],
        rows=storage_rows,
        widths=[16, 28, 18, 23, 20, 25, 21, 18, 18, 18, 18, 20, 18, 18],
        table_name="StokNoktalariTablosu",
        number_formats={column: LITERS_FORMAT for column in range(5, 14)},
        total_row=[
            "GENEL TOPLAM",
            None,
            None,
            None,
            *[sum((row[column] for row in storage_rows), ZERO) for column in range(4, 13)],
            None,
        ],
    )

    daily_summary = list(context["daily_summary"])
    daily_rows = [
        [
            item["day"],
            item["received"] or ZERO,
            item["transferred"] or ZERO,
            item["consumed"] or ZERO,
        ]
        for item in daily_summary
    ]
    weekly_rows = _weekly_consumption_rows(daily_summary)
    daily_sheet = workbook.create_sheet("Günlük Özet")
    daily_sheet.sheet_properties.tabColor = "A17CC1"
    _style_data_sheet(
        daily_sheet,
        headers=["Tarih", "Tedarikçiden gelen", "İç transfer", "Araç tüketimi"],
        rows=daily_rows,
        widths=[16, 22, 18, 20],
        table_name="GunlukOzetTablosu",
        number_formats={1: "dd.mm.yyyy", 2: LITERS_FORMAT, 3: LITERS_FORMAT, 4: LITERS_FORMAT},
        total_row=[
            "GENEL TOPLAM",
            sum((row[1] for row in daily_rows), ZERO),
            sum((row[2] for row in daily_rows), ZERO),
            sum((row[3] for row in daily_rows), ZERO),
        ],
    )
    daily_sheet["G1"] = "Hafta"
    daily_sheet["H1"] = "Araç tüketimi"
    for cell in daily_sheet[1][6:8]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(color="FFFFFF", bold=True, size=10)
    for row_number, (label, total) in enumerate(weekly_rows, start=2):
        daily_sheet.cell(row=row_number, column=7, value=label)
        total_cell = daily_sheet.cell(row=row_number, column=8, value=total)
        total_cell.number_format = LITERS_FORMAT
    daily_sheet.column_dimensions["G"].width = 18
    daily_sheet.column_dimensions["H"].width = 20
    daily_sheet.print_area = f"A1:H{daily_sheet.max_row}"

    _build_summary_sheet(
        summary_sheet,
        context=context,
        generated_by=generated_by,
        generated_at=generated_at,
        region_sheet=region_sheet,
        region_row_count=region_row_count,
        daily_sheet=daily_sheet,
        weekly_row_count=len(weekly_rows),
    )
    workbook.active = 0
    return workbook
